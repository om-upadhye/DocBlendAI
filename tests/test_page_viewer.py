"""PaddleOCR handwriting engine and the document page viewer (lines for scans, blocks for typed files).

The real PaddleOCR engine is never run (conftest refuses it); a fake RapidOCR
result stands in, so these tests stay offline and fast.
"""

from types import SimpleNamespace

import json

import numpy as np
import pytest
from docx import Document as new_docx
from docx.shared import Inches as DocxInches
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.config import settings
from app.modules import format_detection, htr_extractor, ocr_extractor, paddle_extractor, page_store, typed_layout
from tests.conftest import TYPED_PAGE

# --- PaddleOCR engine ----------------------------------------------------------------------


def _quad(x0, y0, x1, y1):
    return np.array([[x0, y0], [x1, y0], [x1, y1], [x0, y1]], dtype=float)


@pytest.fixture
def fake_rapidocr(monkeypatch):
    """Stand in for RapidOCR: returns the given (box, text, score) detections."""

    def _install(detections):
        result = SimpleNamespace(
            boxes=[_quad(*box) for box, _, _ in detections] or None,
            txts=tuple(text for _, text, _ in detections) or None,
            scores=tuple(score for _, _, score in detections) or None,
        )
        monkeypatch.setattr(paddle_extractor, "_engine", lambda threads: (lambda image: result))

    return _install


def test_paddle_groups_boxes_into_lines_in_reading_order(fake_rapidocr) -> None:
    fake_rapidocr(
        [
            ((300, 100, 500, 130), "ability to", 0.9),  # same line, detected out of order
            ((50, 102, 280, 132), "AI system's", 0.8),
            ((50, 200, 400, 232), "is AI planning.", 0.7),
        ]
    )
    lines = []
    text, conf = paddle_extractor.read_image(Image.new("L", (1000, 1000), 255), lines)

    assert text == "AI system's ability to\nis AI planning."
    assert conf == pytest.approx((0.8 * 11 + 0.9 * 10 + 0.7 * 15) / 36)
    assert [line["text"] for line in lines] == ["AI system's ability to", "is AI planning."]
    assert lines[0]["box"] == [0.05, 0.1, 0.5, 0.132]  # union of the two boxes, as page fractions


def test_paddle_blank_page(fake_rapidocr) -> None:
    fake_rapidocr([])
    assert paddle_extractor.read_image(Image.new("L", (100, 100), 255)) == ("", 0.0)


def test_handwriting_engine_is_configurable(monkeypatch, pdf_file) -> None:
    calls = []
    monkeypatch.setattr(paddle_extractor, "paddle_file", lambda *a, **k: calls.append("paddle") or [])
    monkeypatch.setattr(htr_extractor, "htr_file", lambda *a, **k: calls.append("trocr") or [])
    path = str(pdf_file([""]))

    format_detection.handwriting_file(path)
    monkeypatch.setattr(settings, "htr_engine", "trocr")
    format_detection.handwriting_file(path)

    assert calls == ["paddle", "trocr"]


# --- Tesseract line boxes ----------------------------------------------------------------


def test_tesseract_lines_carry_boxes(monkeypatch) -> None:
    data = {
        "text": ["Deep", "learning", "models"], "conf": [90, 80, 70],
        "block_num": [1, 1, 1], "par_num": [1, 1, 1], "line_num": [1, 1, 2],
        "left": [100, 250, 100], "top": [100, 105, 200], "width": [120, 200, 150], "height": [30, 28, 30],
    }
    monkeypatch.setattr(ocr_extractor, "_tesseract_cmd", lambda configured: "tesseract")
    monkeypatch.setattr(ocr_extractor.pytesseract, "image_to_data", lambda image, output_type, timeout: data)

    lines = []
    ocr_extractor.ocr_image(Image.new("L", (1000, 1000), 255), lines)

    assert [line["text"] for line in lines] == ["Deep learning", "models"]
    assert lines[0]["box"] == [0.1, 0.1, 0.45, 0.133]
    assert lines[0]["conf"] == pytest.approx((0.9 * 4 + 0.8 * 8) / 12, abs=1e-4)


# --- viewer endpoints -------------------------------------------------------------------


def _upload(client, path, **form):
    with path.open("rb") as f:
        return client.post("/upload", files={"file": (path.name, f, "application/octet-stream")}, data=form).json()


def test_scanned_page_view_has_image_and_boxed_lines(client, pdf_file, fake_ocr) -> None:
    fake_ocr.text, fake_ocr.conf = "Scanned lecture notes about recognition.", 0.9
    doc = _upload(client, pdf_file(["", ""]))

    pages = client.get(f"/documents/{doc['doc_id']}/pages").json()

    assert [p["page"] for p in pages] == [1, 2]
    first = pages[0]
    assert first["text"] == "Scanned lecture notes about recognition."
    assert first["lines"][0]["box"] == [0.1, 0.1, 0.9, 0.2]
    image = client.get(first["image_url"])
    assert image.status_code == 200 and image.headers["content-type"] == "image/jpeg"


def test_handwritten_page_view(client, pdf_file, fake_ocr, fake_htr) -> None:
    fake_htr.text, fake_htr.conf = "Handwritten notes text", 0.7
    doc = _upload(client, pdf_file([""]), format_hint="handwritten")

    [page] = client.get(f"/documents/{doc['doc_id']}/pages").json()

    assert page["lines"][0]["text"] == "Handwritten notes text"
    assert page["image_url"]


def test_typed_pdf_page_view_has_image_and_boxed_blocks(client, pdf_file) -> None:
    doc = _upload(client, pdf_file([TYPED_PAGE]))

    [page] = client.get(f"/documents/{doc['doc_id']}/pages").json()

    assert "Retrieval-Augmented Generation" in page["text"]
    assert page["lines"] == []
    [block] = page["blocks"]
    assert block["type"] == "paragraph" and "Retrieval-Augmented Generation" in block["text"]
    x0, y0, x1, y1 = block["box"]
    assert 0 < x0 < x1 < 1 and 0 < y0 < y1 < 0.2  # near the top of the page
    assert client.get(page["image_url"]).status_code == 200


def _picture(path) -> str:
    Image.new("RGB", (120, 80), "steelblue").save(path)
    return str(path)


def test_word_page_view_lists_paragraphs_tables_and_pictures(client, tmp_path) -> None:
    path = tmp_path / "notes.docx"
    d = new_docx()
    d.add_heading("Shortest paths", level=1)
    d.add_paragraph("Dijkstra's algorithm finds shortest paths.")
    d.add_paragraph("Uses a priority queue", style="List Bullet")
    d.add_paragraph("Needs non-negative weights", style="List Bullet")
    table = d.add_table(rows=2, cols=2)
    for row, cells in zip(table.rows, [("Algorithm", "Time"), ("Dijkstra", "O(E log V)")]):
        row.cells[0].text, row.cells[1].text = cells
    d.add_picture(_picture(tmp_path / "graph.png"), width=DocxInches(2))
    d.save(path)
    doc = _upload(client, path)

    [page] = client.get(f"/documents/{doc['doc_id']}/pages").json()

    assert page["image_url"] is None
    blocks = page["blocks"]
    assert [b["type"] for b in blocks] == ["paragraph", "paragraph", "paragraph", "table", "image"]
    assert blocks[2]["text"] == "Uses a priority queue\nNeeds non-negative weights"  # one list, one block
    assert blocks[3]["rows"] == [["Algorithm", "Time"], ["Dijkstra", "O(E log V)"]]
    assert all(b["box"] is None for b in blocks)  # Word has no fixed pages to draw boxes on
    picture = client.get(blocks[4]["picture_url"])
    assert picture.status_code == 200 and picture.headers["content-type"] == "image/png"


def test_powerpoint_blocks_follow_each_slide(client, tmp_path) -> None:
    path = tmp_path / "ports.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])  # title only
    slide.shapes.title.text = "Well-known ports"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    for r, (proto, port) in enumerate([("Protocol", "Port"), ("HTTPS", "443")]):
        table.cell(r, 0).text, table.cell(r, 1).text = proto, port
    slide.shapes.add_picture(_picture(tmp_path / "net.png"), Inches(1), Inches(4))
    deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = "Summary"
    deck.save(path)
    doc = _upload(client, path)

    pages = client.get(f"/documents/{doc['doc_id']}/pages").json()

    assert [[b["type"] for b in p["blocks"]] for p in pages] == [["paragraph", "table", "image"], ["paragraph"]]
    assert pages[0]["blocks"][2]["picture_url"]


def test_text_file_blocks_split_on_blank_lines(client, tmp_path) -> None:
    path = tmp_path / "marks.txt"
    path.write_text(
        "Internal assessment marks for the DBMS unit test.\n\n"
        "Name    Test 1    Test 2\nAsha    18    19\nRavi    15    17\n",
        encoding="utf-8",
    )
    doc = _upload(client, path)

    [page] = client.get(f"/documents/{doc['doc_id']}/pages").json()

    assert [b["type"] for b in page["blocks"]] == ["paragraph", "table"]
    assert page["blocks"][1]["rows"][1] == ["Asha", "18", "19"]


def test_typed_document_saved_before_blocks_gets_them_on_view(client, pdf_file) -> None:
    doc = _upload(client, pdf_file([TYPED_PAGE]))
    layout = page_store.pages_dir(doc["doc_id"]) / "layout.json"
    old = [{k: v for k, v in p.items() if k != "blocks"} for p in json.loads(layout.read_text(encoding="utf-8"))]
    layout.write_text(json.dumps(old), encoding="utf-8")  # as saved by version 1.1.0

    [page] = client.get(f"/documents/{doc['doc_id']}/pages").json()

    assert [b["type"] for b in page["blocks"]] == ["paragraph"]
    assert "blocks" in json.loads(layout.read_text(encoding="utf-8"))[0]  # kept for next time


def test_picture_endpoint_serves_only_saved_pictures(client, pdf_file) -> None:
    doc = _upload(client, pdf_file([TYPED_PAGE]))

    for name in ["pic_1.png", "layout.json", "..%2F..%2Fdocblendai.db", "page_1.jpg"]:
        assert client.get(f"/documents/{doc['doc_id']}/pictures/{name}").status_code == 404


def _line(top, bottom, x0=72, x1=540, size=11):
    return {"text": f"line at {top}", "x0": x0, "x1": x1, "top": top, "bottom": bottom, "chars": [{"text": "a", "size": size}]}


def test_pdf_lines_group_into_paragraphs_by_spacing_size_and_column() -> None:
    lines = [
        _line(60, 76, size=16),  # heading: bigger font
        _line(90, 101), _line(105, 116), _line(120, 131),  # paragraph 1, usual 4 pt gap
        _line(140, 151), _line(155, 166),  # paragraph 2 after a 9 pt gap
        _line(90, 101, x0=320, x1=560),  # second column, back at the top
    ]

    groups = typed_layout._group_lines(lines)

    assert [len(g) for g in groups] == [1, 3, 2, 1]


def test_delete_removes_page_view(client, pdf_file, fake_ocr) -> None:
    fake_ocr.text, fake_ocr.conf = "Scanned notes text here.", 0.9
    doc = _upload(client, pdf_file([""]))
    assert page_store.pages_dir(doc["doc_id"]).is_dir()

    client.delete(f"/documents/{doc['doc_id']}")

    assert not page_store.pages_dir(doc["doc_id"]).exists()
    assert client.get(f"/documents/{doc['doc_id']}/pages").status_code == 404


def test_document_uploaded_before_the_viewer_gets_a_clear_404(client, pdf_file) -> None:
    doc = _upload(client, pdf_file([TYPED_PAGE]))
    page_store.delete(doc["doc_id"])  # as if uploaded by an older version

    resp = client.get(f"/documents/{doc['doc_id']}/pages")

    assert resp.status_code == 404 and "upload it again" in resp.json()["detail"]
    assert client.get(f"/documents/{doc['doc_id']}/pages/1/image").status_code == 404
