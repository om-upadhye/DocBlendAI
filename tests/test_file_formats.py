"""Multi-format uploads: images, Word, PowerPoint, and text files (Modules 1-2).

Tesseract, TrOCR, and Gemini are faked (see conftest), so these run offline.
"""

from pathlib import Path

import pytest
from docx import Document as new_docx
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from app.models.schemas import FormatType
from app.modules import file_types, format_detection, ocr_extractor, text_parser, vector_store
from app.modules.ocr_extractor import auto_orient as REAL_AUTO_ORIENT  # bound before conftest's no-op patch
from app.modules.pdf_render import PAGE_LONG_SIDE_INCHES, render_pages
from tests.conftest import fake_embed

# --- builders ---------------------------------------------------------------------


def _write_image(path: Path, size=(800, 600), mode="RGB", **save_kwargs) -> Path:
    color = (255, 255, 255, 0) if mode == "RGBA" else "white"
    img = Image.new(mode, size, color)
    ImageDraw.Draw(img).rectangle((50, 50, 300, 90), fill="black")
    img.save(path, **save_kwargs)
    return path


def _write_docx(path: Path) -> Path:
    doc = new_docx()
    doc.add_heading("Unit 3: Normalization", level=1)
    doc.add_paragraph("Second normal form removes partial dependencies.")
    table = doc.add_table(rows=3, cols=2)
    for row, (name, marks) in zip(table.rows, [("Name", "Marks"), ("Asha", "91"), ("Ravi", "84")]):
        row.cells[0].text, row.cells[1].text = name, marks
    doc.add_paragraph("Third normal form removes transitive dependencies.")
    doc.save(path)
    return path


def _write_pptx(path: Path) -> Path:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])  # title + content
    slide.shapes.title.text = "TCP/IP model"
    slide.placeholders[1].text = "Four layers: link, internet, transport, application"
    slide2 = deck.slides.add_slide(deck.slide_layouts[5])  # title only
    slide2.shapes.title.text = "Ports"
    table = slide2.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(4), Inches(1.5)).table
    for r, (proto, port) in enumerate([("Protocol", "Port"), ("HTTP", "80"), ("HTTPS", "443")]):
        table.cell(r, 0).text, table.cell(r, 1).text = proto, port
    deck.save(path)
    return path


def _upload(client, path: Path, **form):
    with path.open("rb") as f:
        return client.post("/upload", files={"file": (path.name, f, "application/octet-stream")}, data=form)


# --- file kinds -----------------------------------------------------------------------


@pytest.mark.parametrize(
    ("name", "kind"),
    [
        ("notes.PDF", file_types.FileKind.PDF),
        ("photo.JPG", file_types.FileKind.IMAGE),
        ("scan.tiff", file_types.FileKind.IMAGE),
        ("slides.pptx", file_types.FileKind.PPTX),
        ("report.docx", file_types.FileKind.DOCX),
        ("notes.txt", file_types.FileKind.TXT),
        ("old.doc", None),
        ("setup.exe", None),
    ],
)
def test_kind_of(name, kind) -> None:
    assert file_types.kind_of(name) is kind


# --- images ---------------------------------------------------------------------------


def test_multipage_tiff_is_one_page_per_frame(tmp_path) -> None:
    frames = [Image.new("L", (400, 300), 255) for _ in range(3)]
    path = tmp_path / "scan.tiff"
    frames[0].save(path, save_all=True, append_images=frames[1:])

    assert format_detection.page_count(str(path)) == 3
    assert len(list(render_pages(str(path), dpi=150))) == 3
    assert len(list(render_pages(str(path), dpi=150, max_pages=2))) == 2


def test_large_photo_is_scaled_down(tmp_path) -> None:
    path = _write_image(tmp_path / "phone.jpg", size=(3000, 4000))
    [page] = render_pages(str(path), dpi=200)

    assert max(page.size) == round(200 * PAGE_LONG_SIDE_INCHES)
    assert page.mode == "L"


def test_small_screenshot_is_enlarged_at_most_3x(tmp_path) -> None:
    # Regression: a real 576x793 screenshot of notes had ~15 px text, too small for OCR/HTR.
    path = _write_image(tmp_path / "shot.png", size=(576, 793))
    [page] = render_pages(str(path), dpi=200)
    assert max(page.size) == round(200 * PAGE_LONG_SIDE_INCHES)  # 2.95x, under the cap

    tiny = _write_image(tmp_path / "snip.png", size=(300, 100))
    [page] = render_pages(str(tiny), dpi=300)
    assert page.size == (900, 300)  # capped at 3x


def test_phone_photo_is_turned_upright_from_exif(tmp_path) -> None:
    exif = Image.Exif()
    exif[0x0112] = 6  # orientation: rotate 90 degrees clockwise to display
    path = _write_image(tmp_path / "sideways.jpg", size=(800, 600), exif=exif)

    [page] = render_pages(str(path), dpi=150)
    assert page.height > page.width  # stored landscape, displayed portrait
    assert page.width / page.height == pytest.approx(600 / 800, abs=0.01)


def test_transparent_png_gets_white_background(tmp_path) -> None:
    path = _write_image(tmp_path / "clip.png", mode="RGBA")
    [page] = render_pages(str(path), dpi=150)
    sx, sy = page.width / 800, page.height / 600  # the page may be rescaled

    assert page.getpixel((int(700 * sx), int(500 * sy))) == 255  # transparent area is paper, not black
    assert page.getpixel((int(100 * sx), int(70 * sy))) < 50  # the drawn stroke survives


def test_image_upload_goes_through_ocr(client, tmp_path, fake_ocr) -> None:
    fake_ocr.text, fake_ocr.conf = "Photographed lecture notes about scanned text.", 0.9

    resp = _upload(client, _write_image(tmp_path / "board photo.jpg"))

    assert resp.status_code == 201, resp.text
    doc = resp.json()
    assert doc["format_type"] == "scanned"
    assert doc["page_count"] == 1
    assert doc["file_path"].endswith("_board_photo.jpg")  # original name and extension kept


def test_handwritten_photo_goes_through_htr(client, tmp_path, fake_ocr, fake_htr) -> None:
    fake_ocr.text, fake_ocr.conf = "sc ribb le", 0.2
    fake_htr.text, fake_htr.conf = "Handwritten notes text", 0.8

    resp = _upload(client, _write_image(tmp_path / "notes.png"))

    assert resp.status_code == 201
    assert resp.json()["format_type"] == "handwritten"


def test_typed_hint_is_ignored_for_images(client, tmp_path, fake_ocr) -> None:
    fake_ocr.text, fake_ocr.conf = "Photographed notes about scanned text.", 0.9
    resp = _upload(client, _write_image(tmp_path / "photo.png"), format_hint="typed")
    assert resp.json()["format_type"] == "scanned"  # an image has no text layer to parse


def test_corrupt_image_is_422_and_cleaned_up(client, tmp_path) -> None:
    bad = tmp_path / "broken.png"
    bad.write_bytes(b"\x89PNG\r\n\x1a\nnot really a png")

    resp = _upload(client, bad)

    assert resp.status_code == 422
    assert ".png" in resp.json()["detail"]
    assert not any((tmp_path / "uploads").iterdir())


# --- Word -------------------------------------------------------------------------------


def test_docx_keeps_paragraphs_and_tables_in_order(tmp_path) -> None:
    [(text, conf)] = text_parser.parse(str(_write_docx(tmp_path / "unit3.docx")))

    assert conf == 1.0
    lines = text.splitlines()
    assert lines[0] == "Unit 3: Normalization"
    assert "Name  Marks" in lines and "Asha  91" in lines  # cells kept apart, like a table
    assert lines.index("Asha  91") < lines.index("Third normal form removes transitive dependencies.")


def test_docx_upload_is_typed_even_with_a_scanned_hint(client, tmp_path) -> None:
    resp = _upload(client, _write_docx(tmp_path / "unit3.docx"), format_hint="scanned")

    assert resp.status_code == 201, resp.text
    assert resp.json()["format_type"] == "typed"
    assert resp.json()["page_count"] == 1
    hits = vector_store.search(fake_embed(["second normal form partial dependencies"], "")[0], top_k=5)
    assert any("partial dependencies" in c.text for c, _ in hits)


def test_corrupt_docx_is_422(client, tmp_path) -> None:
    bad = tmp_path / "broken.docx"
    bad.write_bytes(b"PK\x03\x04 not a real zip")
    assert _upload(client, bad).status_code == 422


# --- PowerPoint -------------------------------------------------------------------------


def test_pptx_is_one_page_per_slide_with_tables(tmp_path) -> None:
    pages = text_parser.parse(str(_write_pptx(tmp_path / "networks.pptx")))

    assert len(pages) == 2
    assert "TCP/IP model" in pages[0][0] and "Four layers" in pages[0][0]
    assert "HTTPS  443" in pages[1][0].splitlines()


def test_pptx_upload(client, tmp_path) -> None:
    resp = _upload(client, _write_pptx(tmp_path / "networks.pptx"))
    assert resp.status_code == 201
    assert resp.json()["format_type"] == "typed"
    assert resp.json()["page_count"] == 2


# --- plain text -------------------------------------------------------------------------


def test_txt_utf8_with_bom(tmp_path) -> None:
    path = tmp_path / "notes.txt"
    path.write_bytes("﻿Schrödinger equation — notes".encode("utf-8"))
    assert text_parser.parse(str(path)) == [("Schrödinger equation — notes", 1.0)]


def test_txt_windows_encoding_falls_back(tmp_path) -> None:
    path = tmp_path / "notes.txt"
    path.write_bytes("café résumé".encode("cp1252"))
    assert text_parser.parse(str(path)) == [("café résumé", 1.0)]


def test_binary_file_named_txt_is_422(client, tmp_path) -> None:
    bad = tmp_path / "data.txt"
    bad.write_bytes(b"\x00\x01\x02 binary")
    assert _upload(client, bad).status_code == 422


def test_txt_upload(client, tmp_path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("Dijkstra's algorithm finds shortest paths in weighted graphs.", encoding="utf-8")

    resp = _upload(client, path)

    assert resp.status_code == 201
    assert resp.json()["format_type"] == "typed"


# --- orientation ------------------------------------------------------------------------
# The conftest replaces ocr_extractor.auto_orient with a no-op; REAL_AUTO_ORIENT is the original.


def _page_with_marker(size=(600, 400)) -> Image.Image:
    page = Image.new("L", size, 255)
    ImageDraw.Draw(page).rectangle((10, 10, 60, 30), fill=0)  # marker in the top-left corner
    return page


@pytest.fixture
def fake_osd(monkeypatch):
    """Fake Tesseract OSD: suggests `rotate`; OCR confidence is high only for upright pages."""

    def _install(rotate: int, upright_test):
        monkeypatch.setattr(ocr_extractor, "_tesseract_cmd", lambda configured: "tesseract")
        monkeypatch.setattr(
            ocr_extractor.pytesseract, "image_to_osd", lambda image, output_type, timeout: {"rotate": rotate}
        )
        monkeypatch.setattr(ocr_extractor, "_quick_conf", lambda img: 0.9 if upright_test(img) else 0.2)

    return _install


def test_upright_page_is_left_alone(fake_osd) -> None:
    page = _page_with_marker()
    fake_osd(0, lambda img: True)
    assert REAL_AUTO_ORIENT(page) is page


def test_upside_down_page_is_turned_upright(fake_osd) -> None:
    upright = _page_with_marker()
    upside_down = upright.rotate(180)
    fake_osd(180, lambda img: img.getpixel((30, 20)) == 0)  # upright iff the marker is top-left

    fixed = REAL_AUTO_ORIENT(upside_down)
    assert fixed.getpixel((30, 20)) == 0


def test_sideways_page_is_turned_upright_whichever_way_osd_means(fake_osd) -> None:
    upright = _page_with_marker(size=(600, 400))
    sideways = upright.rotate(90, expand=True)  # now 400 x 600
    fake_osd(270, lambda img: img.size == (600, 400) and img.getpixel((30, 20)) == 0)

    fixed = REAL_AUTO_ORIENT(sideways)
    assert fixed.size == (600, 400) and fixed.getpixel((30, 20)) == 0


def test_wrong_osd_suggestion_is_rejected(fake_osd) -> None:
    page = _page_with_marker()
    fake_osd(180, lambda img: img.getpixel((30, 20)) == 0)  # OSD is wrong: the page is already upright
    assert REAL_AUTO_ORIENT(page) is page


def test_osd_failure_leaves_page_unchanged(monkeypatch) -> None:
    def _fail(image, output_type, timeout):
        raise ocr_extractor.pytesseract.TesseractError(1, "Too few characters. Skipping this page")

    monkeypatch.setattr(ocr_extractor, "_tesseract_cmd", lambda configured: "tesseract")
    monkeypatch.setattr(ocr_extractor.pytesseract, "image_to_osd", _fail)
    page = _page_with_marker()
    assert REAL_AUTO_ORIENT(page) is page


def test_resolve_format_rules(tmp_path, fake_ocr) -> None:
    fake_ocr.text, fake_ocr.conf = "Printed words", 0.9
    txt = tmp_path / "a.txt"
    txt.write_text("hello", encoding="utf-8")
    img = _write_image(tmp_path / "a.png")

    assert format_detection.resolve_format(str(txt), FormatType.HANDWRITTEN) is FormatType.TYPED
    assert format_detection.resolve_format(str(img), FormatType.HANDWRITTEN) is FormatType.HANDWRITTEN
    assert format_detection.resolve_format(str(img), FormatType.TYPED) is FormatType.SCANNED
    assert format_detection.resolve_format(str(img), None) is FormatType.SCANNED
