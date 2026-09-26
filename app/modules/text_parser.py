"""Module 2 — direct parsing path (typed documents).

Responsibility: read text that is already digital, with no recognition step,
so raw_conf is 1.0 for every page:
- PDF: the embedded text layer (pdfplumber), one entry per page
- DOCX: paragraphs and tables in document order, as one entry (Word files
  have no fixed pages until rendered)
- PPTX: one entry per slide (text boxes, tables, grouped shapes)
- TXT: the whole file as one entry

Table rows are written with cells separated by two spaces, the same shape
content_type.py (Module 4) recognizes as a table.

Uses from schemas.py: nothing directly; output feeds chunker.py -> RecognizedChunk.
"""

import zipfile

import pdfplumber
from docx import Document as open_docx
from docx.opc.exceptions import PackageNotFoundError as DocxPackageError
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PackageNotFoundError as PptxPackageError

from app.modules.file_types import FileKind, UnreadableFileError, kind_of

TYPED_CONFIDENCE = 1.0
CELL_SEPARATOR = "  "

_BROKEN_PACKAGE = (DocxPackageError, PptxPackageError, zipfile.BadZipFile, KeyError, ValueError)


def parse(file_path: str) -> list[tuple[str, float]]:
    """Return (page_text, raw_conf=1.0) per page for any typed file kind."""
    kind = kind_of(file_path)
    if kind is FileKind.DOCX:
        return parse_docx(file_path)
    if kind is FileKind.PPTX:
        return parse_pptx(file_path)
    if kind is FileKind.TXT:
        return parse_txt(file_path)
    return parse_pdf(file_path)


def parse_pdf(file_path: str) -> list[tuple[str, float]]:
    """Return (page_text, raw_conf=1.0) for each page; pages without text give ""."""
    with pdfplumber.open(file_path) as pdf:
        return [((page.extract_text() or "").strip(), TYPED_CONFIDENCE) for page in pdf.pages]


def _row_text(cells: list[str]) -> str:
    # Merged cells repeat the same text in every cell they span; keep it once.
    deduped = [c for i, c in enumerate(cells) if c and (i == 0 or c != cells[i - 1])]
    return CELL_SEPARATOR.join(deduped)


def _table_lines(rows) -> list[str]:
    lines = []
    for row in rows:
        line = _row_text([cell.text.strip().replace("\n", " ") for cell in row.cells])
        if line:
            lines.append(line)
    return lines


def parse_docx(file_path: str) -> list[tuple[str, float]]:
    try:
        doc = open_docx(file_path)
    except _BROKEN_PACKAGE as e:
        raise UnreadableFileError(f"not a readable Word document: {e}") from e

    lines: list[str] = []
    # Walk the body in order so tables stay where they appear between paragraphs.
    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph):
            if block.text.strip():
                lines.append(block.text.strip())
        elif isinstance(block, DocxTable):
            lines.extend(_table_lines(block.rows))
    return [("\n".join(lines), TYPED_CONFIDENCE)]


def _shape_lines(shape) -> list[str]:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [line for inner in shape.shapes for line in _shape_lines(inner)]
    if getattr(shape, "has_table", False) and shape.has_table:
        return _table_lines(shape.table.rows)
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        return [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
    return []


def parse_pptx(file_path: str) -> list[tuple[str, float]]:
    try:
        deck = Presentation(file_path)
    except _BROKEN_PACKAGE as e:
        raise UnreadableFileError(f"not a readable PowerPoint file: {e}") from e
    return [
        ("\n".join(line for shape in slide.shapes for line in _shape_lines(shape)), TYPED_CONFIDENCE)
        for slide in deck.slides
    ]


def parse_txt(file_path: str) -> list[tuple[str, float]]:
    raw = open(file_path, "rb").read()
    if b"\x00" in raw[:4096]:
        raise UnreadableFileError("not a text file (binary content)")
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252", errors="replace")  # common for Notepad files saved on Windows
    return [(text.strip(), TYPED_CONFIDENCE)]
